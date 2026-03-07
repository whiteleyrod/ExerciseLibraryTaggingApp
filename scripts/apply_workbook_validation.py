from __future__ import annotations

from pathlib import Path
from typing import Iterable

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / 'Exercise Library Index.xlsx'
HELPER_SHEET = 'ValidationLists'
MUSCLE_LIST_NAME = 'MuscleAreaList'
SOURCE_SHEET = 'TaggingCategories'
SOURCE_HEADER = 'Muscle area'


def find_header_column(sheet, header: str, scan_rows: int = 5) -> tuple[int, int] | None:
    for row in range(1, scan_rows + 1):
        for col in range(1, sheet.max_column + 1):
            value = str(sheet.cell(row=row, column=col).value or '').strip()
            if value == header:
                return col, row
    return None


def unique_non_empty(values: Iterable[str]) -> list[str]:
    return sorted({value.strip() for value in values if value and value.strip()})


def remove_existing_named_range(workbook, name: str) -> None:
    to_remove = [item for item in workbook.defined_names.values() if item.name == name]
    for item in to_remove:
        workbook.defined_names.delete(item.name)


def apply_validation() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f'Workbook not found: {WORKBOOK_PATH}')

    workbook = openpyxl.load_workbook(WORKBOOK_PATH)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise ValueError(f'Missing source sheet: {SOURCE_SHEET}')

    source_ws = workbook[SOURCE_SHEET]
    source_header = find_header_column(source_ws, SOURCE_HEADER)
    if not source_header:
        raise ValueError(f'Could not find "{SOURCE_HEADER}" column in {SOURCE_SHEET}')

    source_col, source_header_row = source_header
    source_values = [
        str(source_ws.cell(row=row, column=source_col).value or '').strip()
        for row in range(source_header_row + 1, source_ws.max_row + 1)
    ]
    muscle_areas = unique_non_empty(source_values)
    if not muscle_areas:
        raise ValueError('No muscle area values found in TaggingCategories sheet')

    if HELPER_SHEET in workbook.sheetnames:
        helper_ws = workbook[HELPER_SHEET]
        helper_ws.delete_cols(1, helper_ws.max_column or 1)
    else:
        helper_ws = workbook.create_sheet(HELPER_SHEET)

    for idx, value in enumerate(muscle_areas, start=1):
        helper_ws.cell(row=idx, column=1, value=value)

    helper_ws.sheet_state = 'hidden'

    remove_existing_named_range(workbook, MUSCLE_LIST_NAME)
    reference = f"'{HELPER_SHEET}'!$A$1:$A${len(muscle_areas)}"
    workbook.defined_names.add(DefinedName(name=MUSCLE_LIST_NAME, attr_text=reference))

    target_sheets = [
        sheet_name
        for sheet_name in workbook.sheetnames
        if sheet_name not in {SOURCE_SHEET, 'ExerciseName&Link', HELPER_SHEET}
    ]

    applied = []
    for sheet_name in target_sheets:
        ws = workbook[sheet_name]
        header = find_header_column(ws, SOURCE_HEADER)
        if not header:
            continue

        col_index, header_row = header
        col_letter = get_column_letter(col_index)
        start_row = header_row + 1
        end_row = max(ws.max_row, start_row)

        existing = list(ws.data_validations.dataValidation)
        ws.data_validations.dataValidation = [dv for dv in existing if dv.formula1 != f'={MUSCLE_LIST_NAME}']

        data_validation = DataValidation(type='list', formula1=f'={MUSCLE_LIST_NAME}', allow_blank=True)
        data_validation.error = 'Select a value from the approved Muscle area list.'
        data_validation.errorTitle = 'Invalid muscle area'
        data_validation.prompt = 'Choose muscle area from dropdown options.'
        data_validation.promptTitle = 'Muscle area'

        ws.add_data_validation(data_validation)
        data_validation.add(f'{col_letter}{start_row}:{col_letter}{end_row}')
        applied.append((sheet_name, f'{col_letter}{start_row}:{col_letter}{end_row}'))

    workbook.save(WORKBOOK_PATH)

    print(f'Applied {MUSCLE_LIST_NAME} with {len(muscle_areas)} values.')
    if not applied:
        print('No target sheets with a Muscle area column were found.')
    else:
        print('Applied to:')
        for sheet_name, cell_range in applied:
            print(f'  - {sheet_name}: {cell_range}')


if __name__ == '__main__':
    apply_validation()
