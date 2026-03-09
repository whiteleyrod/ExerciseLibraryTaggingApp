from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / 'Muscle, joint, area.xlsx'
EQUIPMENT_WORKBOOK = ROOT / 'EquipmentTags.xlsx'
PUBLIC_DIR = ROOT / 'web' / 'public'

PRIMARY_SHEET = 'musclejointarea'
GROUP_SHEETS = ['upperbody', 'trunk&core', 'lowerbody']
EQUIPMENT_HEADERS = ['Category', 'Equipment Included']


def normalize(value: object) -> str:
    if value is None:
        return ''
    return str(value).strip()


def read_sheet_rows(workbook: openpyxl.Workbook, sheet_name: str) -> list[dict[str, str]]:
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [normalize(cell) for cell in rows[0]]
    output: list[dict[str, str]] = []
    for row in rows[1:]:
        record: dict[str, str] = {}
        has_content = False
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = normalize(row[idx] if idx < len(row) else '')
            if value:
                has_content = True
            record[header] = value
        if has_content:
            output.append(record)
    return output


def write_csv(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: row.get(header, '') for header in headers})


def write_json(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding='utf-8')


def main() -> None:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True)

    primary_rows = read_sheet_rows(workbook, PRIMARY_SHEET)
    primary_headers = ['Muscles involved', 'Joints involved', 'Muscle area']
    write_csv(PUBLIC_DIR / 'MuscleJointArea.csv', primary_headers, primary_rows)
    write_json(PUBLIC_DIR / 'MuscleJointArea.json', primary_rows)

    grouped_rows: list[dict[str, str]] = []
    grouped_headers = ['Group', 'Muscle Area Tag', 'Muscles Involved', 'Joints Involved']

    for sheet_name in GROUP_SHEETS:
        rows = read_sheet_rows(workbook, sheet_name)
        for row in rows:
            grouped_rows.append(
                {
                    'Group': sheet_name,
                    'Muscle Area Tag': row.get('Muscle Area Tag', ''),
                    'Muscles Involved': row.get('Muscles Involved', ''),
                    'Joints Involved': row.get('Joints Involved', ''),
                }
            )

    write_csv(PUBLIC_DIR / 'MusclePaneGroups.csv', grouped_headers, grouped_rows)
    write_json(PUBLIC_DIR / 'MusclePaneGroups.json', grouped_rows)

    equipment_workbook = openpyxl.load_workbook(EQUIPMENT_WORKBOOK, data_only=True)
    equipment_sheet = equipment_workbook[equipment_workbook.sheetnames[0]]
    equipment_rows = read_sheet_rows(equipment_workbook, equipment_sheet.title)

    normalized_equipment_rows: list[dict[str, str]] = []
    for row in equipment_rows:
        normalized_equipment_rows.append(
            {
                'Category': row.get('Category', ''),
                'Equipment Included': row.get('Equipment Included', ''),
            }
        )

    write_csv(PUBLIC_DIR / 'EquipmentTags.csv', EQUIPMENT_HEADERS, normalized_equipment_rows)
    write_json(PUBLIC_DIR / 'EquipmentTags.json', normalized_equipment_rows)
    print(f'Wrote {len(primary_rows)} rows to web/public/MuscleJointArea.csv')
    print(f'Wrote {len(primary_rows)} rows to web/public/MuscleJointArea.json')
    print(f'Wrote {len(grouped_rows)} rows to web/public/MusclePaneGroups.csv')
    print(f'Wrote {len(grouped_rows)} rows to web/public/MusclePaneGroups.json')
    print(f'Wrote {len(normalized_equipment_rows)} rows to web/public/EquipmentTags.csv')
    print(f'Wrote {len(normalized_equipment_rows)} rows to web/public/EquipmentTags.json')


if __name__ == '__main__':
    main()
